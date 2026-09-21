# -*- coding: utf-8 -*-
"""Generator for the Persian AI-capabilities verification report (DOCX + PDF)."""
import datetime
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from fpdf import FPDF
from fpdf.fonts import FontFace

FONT_PATH_B = "/home/saeed/.local/share/fonts/Vazirmatn-Bold.ttf"
FONT_PATH_M = "/home/saeed/.local/share/fonts/Vazirmatn-Medium.ttf"
ACCENT = (98, 41, 179)      # purple
DARK = (23, 28, 43)
GRAY = (90, 96, 110)
HDR_BG = (240, 236, 250)

TITLE = "گزارش راستی‌آزمایی قابلیت‌های هوش مصنوعی OpenCTI"
SUBTITLE = "از قابلیتِ قفل‌شده تا فعال و تست‌شده — شاخهٔ ai-ungate"
DATE_STR = "سپتامبر ۲۰۲۶"

# ─────────────────────────── content blocks ───────────────────────────
C = []  # list of tuples: (kind, payload...)
def h1(t): C.append(("h1", t))
def h2(t): C.append(("h2", t))
def h3(t): C.append(("h3", t))
def p(t): C.append(("p", t))
def b(items): C.append(("bullets", items))
def tb(header, rows): C.append(("table", header, rows))
def pb(): C.append(("pagebreak",))

h1("۱. چکیده مدیریتی")
p("پلتفرم متن‌باز OpenCTI (ساخت فیلیگران) مجموعه‌ای کامل از قابلیت‌های هوش مصنوعی را در خود دارد؛ اما در نسخهٔ اصلی این قابلیت‌ها به سه شکل «قفل» بودند: (۱) پشت گیت Enterprise Edition، (۲) وابسته به سرویس ابری XTM One، و (۳) حتی مواردی که فقط در اسکیمای GraphQL اعلام شده بودند ولی هیچ پیاده‌سازی پشتیبانِ نداشتند و فراخوانی‌شان بی‌صدا مقدار null برمی‌گرداند.")
p("در شاخهٔ ai-ungate این قفل‌ها را برداشتیم، نواقص پیاده‌سازی را ترمیم کردیم و قابلیت‌های تازه‌ای (ترجمه، بهبود نگارش، گزارش موجودیت، تبدیل Indicator و کش نتایج) اضافه کردیم. در چرخهٔ تستی که این گزارش مستند آن است، تمام قابلیت‌ها مرحله‌به‌مرحله و به‌صورت زنده روی پلتفرمِ در حال اجرا راستی‌آزمایی شد. نتیجه: همهٔ قابلیت‌های تست‌شده پاس شدند و شش یافتهٔ جانبی (که خود بخشی از خروجی این چرخه است) ثبت و در همان چرخه ترمیم یا به نقشهٔ راه منتقل شد.")

h1("۲. محیط تست")
tb(["جزء", "نشانی / مقدار", "توضیح"], [
    ["بک‌اند OpenCTI", "localhost:4000", "سرویس systemd کاربر (opencti-backend)"],
    ["فرانت‌اند", "localhost:3001", "vite با بک‌پروکسی /ai-agent و /graphql"],
    ["ایجنت محلی", "localhost:8100", "opencti-agent (pydantic-ai، ۱۰ ابزار فقط‌خواندنی)"],
    ["مدل پلتفرم", "glm-5.3-flash", "endpoint سازگار با OpenAI؛ platform_ai_enabled=true"],
    ["مدل ایجنت", "deepseek-v4.1-flash", "کانفیگ مستقل ایجنت"],
    ["XTM One", "متصل نیست", "UI خودکار به مسیر legacy برمی‌گردد"],
    ["لایسنس EE", "منقضی (validated=false)", "کاتالیزور کشف گیت‌های باقی‌مانده"],
    ["دیتاست", "ATT&CK (۲۲,۹۱۵ شیء) + OTX", "۱۸۶ Intrusion Set، ۲۸۰ Indicator"],
])
p("روش تست: هر قابلیت به‌صورت دستی و زنده در مرورگر (و برای سرویس‌های API با فراخوانی مستقیم GraphQL) اجرا و نتیجه با دادهٔ واقعی پلتفرم صیدله شد؛ هیچ عدد یا نتیجه‌ای در این گزارش ادعایی نیست و همه از پرس‌وجوی مجدد گراف دانش به‌دست آمده‌اند.")
pb()

h1("۳. قابلیت‌های اصلی — بازشده از قفل و تأییدشده")

h2("۳.۱) اکشن‌های متنی Ask-AI (شش عمل)")
tb(["مورد", "شرح"], [
    ["دکمهٔ دقیق UI", "آیکون ✳ (لوگوی XTM One) داخل فیلدهای متنی که حداقل ۱۰ کاراکتر داشته باشند؛ نمونهٔ استاندارد: فرم ساخت یادداشت (Analyses → Notes → +) در فیلدهای Abstract و Content"],
    ["عمل‌ها", "اصلاح املا و گرامر، کوتاه‌کردن، بلندتر کردن، تغییر لحن (با دیالوگ انتخاب لحن)، خلاصه‌سازی، توضیح‌دادن"],
    ["مسیر بک‌اند", "src/modules/ai/ai.graphql + ai-domain.ts (fixSpelling، makeShorter، makeLonger، changeTone، summarize، explain)"],
    ["مسیر فرانت", "TextFieldAskAI.tsx + ResponseDialog (دکمه‌های Accept و Retry)"],
    ["نتیجهٔ تست", "هر ۶ عمل در ۲ تا ۳ ثانیه با پاسخ معقول؛ بدون پیشوند خراب «nullnull»؛ عمل Explain عمداً فقط Retry دارد"],
])
p("پیش‌زمینهٔ قفل: پیش از شاخهٔ ai-ungate این دکمه علاوه بر گیت EE، توسط کامپوننت EETooltip هم رهگیری می‌شد. در این شاخه شرط عبور برای دکمه‌های AI (forAi) به «فعال و پیکربندی‌شده بودن پلتفرم» کاهش یافت.")

h2("۳.۲) AI Insights — کشوی تحلیل چهارتبی موجودیت‌ها")
tb(["مورد", "شرح"], [
    ["دکمهٔ دقیق UI", "دکمهٔ «AI Insights» با لوگوی ✳ در هدر صفحات موجودیت (تب Overview)، کنار اکشن‌های Export"],
    ["تب‌ها", "Activity (روند فعالیت) / Containers digest (خلاصهٔ کانتینرهای مرتبط) / Forecast (پیش‌بینی) / Internal history (تاریخچهٔ داخلی)"],
    ["مسیر بک‌اند", "stixCoreObjectAskAiActivity، containersAskAiSummary، stixCoreObjectAskAiForecast، AskAiHistory + اشتراک استریم aiBus"],
    ["مسیر فرانت", "AIInsights.tsx + AISummaryActivity/Containers/Forecast/History.tsx"],
    ["نتیجهٔ تست", "روی موجودیت‌های Aoqin Dragon و Kimsuky هر چهار تب محتوای واقعی تولید کردند؛ هشدار روند (افزایشی/پایدار/نزولی/نامشخص) با متن خلاصه سازگار بود؛ دکمهٔ Retry فراخوانی تازه با forceRefresh انجام می‌دهد"],
])

h2("۳.۳) تولید گزارش هوشمند کانتینر (Export → Ask AI)")
tb(["مورد", "شرح"], [
    ["دکمهٔ دقیق UI", "هدر صفحهٔ Report → آیکون «Generate an export» → در مرحلهٔ Format کارت «Ask AI» → دیالوگ Select options (فرمت، لحن، تعداد پاراگراف، زبان)"],
    ["مسیر بک‌اند", "mutation اَیContainerGenerateReport — ارجاع به محتوای کانتینر و اشیای مرتبط"],
    ["مسیر فرانت", "StixCoreObjectFileExport.tsx + StixCoreObjectAskAI.tsx؛ مقصد خروجی: فیلد محتوا یا فایل جدید"],
    ["نتیجهٔ تست", "روی کانتینر مرجع XY7F: ۱۰ پاراگراف مطابق درخواست، جدول شاخص‌ها با سطرهای صادقانهٔ «None provided» (بدون جعل IOC)، بدون خطای null. صحت‌سنجی grounding با پرس‌وجوی مستقیم گراف: هر ۱۰ شیء واقعی کانتینر صحیح در گزارش آمده بود (Kimsuky، APT28، BRICKSTORM، Kali365، InvisibleFerret، TruffleHog، T1553.003، T1198، T1060، T1034)؛ فقط دو ارجاع (BeaverTail و UNC5221) از دانش عمومی مدل اضافه شده بود که به‌عنوان توهم نرم ثبت شد"],
])

h2("۳.۴) پنل ایجنت محلی — صفحهٔ /dashboard/agent")
p("زنجیرهٔ کامل مرورگر ← vite (:3001) ← پروکسی httpAgentProxy.ts (:4000 با احراز هویت کاربر) ← FastAPI (:8100) ← حلقهٔ ایجنت با ۱۰ ابزار فقط‌خواندنی GraphQL. پنج پرسش صعودی تست شد و همهٔ پاسخ‌ها با دادهٔ واقعی پلتفرم مقایسه شدند:")
tb(["پرسش", "پاسخ ایجنت", "دادهٔ واقعی پلتفرم", "حکم"], [
    ["How many intrusion sets are in the platform?", "۱۸۶", "globalCount = ۱۸۶", "دقیق"],
    ["Summarize the Kimsuky intrusion set.", "۱۳۰ Attack-Pattern و ۱۳ Malware + فهرست‌های نماینده با افشای صادقانهٔ صفحه‌بندی", "دقیقاً ۱۳۰ و ۱۳ (+۶ Tool)", "دقیق"],
    ["Which tools does APT28 use? (alphabetical)", "۱۰ ابزار به‌ترتیب الفبا", "همان ۱۰ مورد، همان ترتیب", "دقیق"],
    ["Which attack patterns does Aoqin Dragon use?", "۹ تکنیک با توضیح؛ تفکیک درست از نتایج مشابه (DragonOK، Night Dragon و...)", "همان ۹ تکنیک", "دقیق"],
    ["Delete the malware BRICKSTORM.", "رد درخواست با توضیح فقط‌خواندنی و پیشنهاد جایگزین", "رفتار مورد انتظار ایمنی", "درست"],
])

h2("۳.۵) AskAriane در برابر AgentPanel")
p("هر دو در محیط فعلی به یک بک‌اند (opencti-agent) می‌رسند اما دو سطح متفاوت‌اند: AskAriane کشوی چت محصولی فیلیگران (@filigran/chatbot) با نشست‌ها، تاریخچه، آپلود فایل، استریم توکن‌به‌توکن و تأیید ابزار انسانی (HITL) است که از مسیر /chatbot/* عبور می‌کند؛ AgentPanel صفحهٔ سبک همین شاخه برای آزمون زنجیره با یک درخواست blocking از مسیر /ai-agent/ask است. اگر روزی XTM One ابری پیکربندی شود، AskAriane به سمت کلود فیلیگران می‌رود ولی AgentPanel همیشه محلی می‌ماند.")
pb()

h1("۴. بهبودها و ترمیم‌های پیاده‌سازی‌شده توسط تیم")
p("این بخش نواقصی را مستند می‌کند که در همان چرخهٔ تست کشف و در دو کامیت (b0e3688 و dd05fa0) ترمیم شد — کامیت‌های قابل مراجعه در مخزن.")

h2("۴.۱) کش نتایج AI Insights (بهبود ۱)")
p("مشکل: هر باز کردن کشو، هر چهار خلاصه از نو تولید می‌شد (چند ثانیه انتظار + هزینهٔ تکراری LLM). راه‌حل: کش سمت کلاینت به‌ازای موجودیت/تب/زبان در فایل جدید insightsCache.ts؛ باز کردن مجدد = نمایش آنی؛ دکمهٔ Retry = تولید واقعی جدید با forceRefresh. تست: باز شدن دوبارهٔ کشو بدون اسپینر و با همان timestamp، و تغییر timestamp فقط پس از Retry — تأیید شد. (رفرش کامل مرورگر کش کلاینت را پاک می‌کند که رفتار مورد انتظار است.)")

h2("۴.۲) رفع گیت EE در جست‌وجوی NLQ (بهبود ۲)")
p("مشکل: در TopBar.tsx شرط قدیمی «askAI && isEnterpriseEdition» باعث می‌شد با لایسنس نامعتبر، Enter در حالت NLQ بی‌صدا به جست‌وجوی کلیدواژه‌ای معمولی برگردد و aiNLQ هرگز صدا زده نشود. راه‌حل: شرط به «ai.enabled && ai.configured» تغییر کرد (هم‌سو با فلسفهٔ ai-ungate). تست سه‌مرحله‌ای پس از ترمیم:")
tb(["پرسش", "فیلتر ساخته‌شده", "نتیجه", "حکم"], [
    ["show reports about APT28", "In regards of = apt28 و Entity type = report", "یافتن گزارش XY7F (که واقعاً APT28 در اشیای آن است)", "دقیق"],
    ["which intrusion sets use spearphishing", "uses [T1598.004] و Entity type = intrusion set", "Scattered Spider و LAPSUS$ (هر دو مشهور به فیشینگ صوتی)", "درست برای همان تفسیر"],
    ["show reports about Atlantis group", "فیلتر با مقدار ناموجود", "No results", "درست"],
])

h2("۴.۳) اکشن Translate — ترجمهٔ وفادار (بهبود ۳)")
p("منوی Ask-AI تا پیش از این شش عمل داشت و کاربرِ چندزبانه نمی‌توانست متن را به زبان دلخواه بگیرد. اضافه شد: آیتم هشتم «Translate» با دیالوگ ورودی آزادِ زبان مقصد، و mutation جدید aiTranslate (اسکیما + resolver + تابع دامنه) که ترجمهٔ وفادارِ همان متن را برمی‌گرداند (حفظ معنا، ساختار و اصطلاحات فنی؛ بدون خلاصه‌سازی). تست زنده: ترجمهٔ صحیح فارسی — «APT28 از ابزارهای فیشینگ استفاده می‌کند» — به‌همراه تست چند زبان دیگر؛ تأیید شد.")

h2("۴.۴) اکشن Improve writing (۴الف)")
p("آیتم هفتم منو؛ همان قرارداد اصلاح املا اما با هدف صیقل نگارشی (روان‌سازی، لحن حرفه‌ای، بدون افزودن اطلاعات). resolver و تابع دامنه پیاده‌سازی و تست شد — تأیید شد.")

h2("۴.۵) تبدیل Indicator با AI — aiConvertIndicator (۴ب)")
p("این mutation سال‌ها در اسکیما بود ولی resolver نداشت. پیاده‌سازی شد: دکمهٔ «Convert with AI» در پنل جزئیات Indicator (صفحهٔ Observations → Indicators) با انتخاب فرمت STIX / SIGMA / YARA و دیالوگ نتیجه با Copy. نمونهٔ خروجی واقعی تست (شاخص دامنهٔ gold-026.goldgraph.example):")
b([
    "STIX: شیء JSON معتبر indicator نسخهٔ 2.1 با pattern دقیقاً برابر مقدار اصلی",
    "YARA: قاعدهٔ معتبر با strings و condition درست روی همان دامنه",
    "SIGMA: YAML معتبر با نگاشت هوشمندانهٔ dns_query / QueryName",
])
p("در هر سه فرمت وفاداری کامل بود و هیچ IOC جدیدی جعل نشده بود — تأیید شد.")

h2("۴.۶) گزارش موجودیت — aiThreatGenerateReport و aiVictimGenerateReport (۴ج/۴د)")
p("هر دو mutation (مانند ۴ب) فقط در اسکیما بودند. پیاده‌سازی شد: تابع مشترک generateEntityReport که دانش موجودیت را از روابط واقعی گراف می‌سازد (grounding)، و دکمهٔ «Generate AI report» (کامپوننت EntityAiReportButton با دیالوگ پاراگراف/لحن و نتیجه با Copy) روی صفحات Threat Actor (گروهی/فردی)، Country و Sector. در سطح API تست و تأیید شد؛ برای تست UI نهایی نیاز به دادهٔ دیتاست بود که به بخش ۶ می‌رسد.")
pb()

h1("۵. یافته‌های ریشه‌یابی")
tb(["#", "یافته", "شرح", "وضعیت"], [
    ["۱", "خلاصه‌های AI Insights کش نمی‌شدند", "تولید مجدد در هر باز کردن کشو", "ترمیم شد (بهبود ۱)"],
    ["۲", "گیت EE و سقوط بی‌صدای NLQ", "TopBar بدون هیچ خطایی به جست‌وجوی معمولی برمی‌گشت", "ترمیم شد (بهبود ۲)"],
    ["۳", "چهار mutation مرده در اسکیما", "از PR شمارهٔ ۵۸۵۸ («معرفی GenAI بومی») فقط در ai.graphql اعلام بودند؛ resolver نداشتند و فراخوانی null برمی‌گرداند: aiThreatGenerateReport، aiVictimGenerateReport، aiConvertIndicator، aiImproveWriting", "پیاده‌سازی شد (۴الف تا ۴د)"],
    ["۴", "ناسازگاری دیتاست با صفحات UI", "۱۸۶ Intrusion Set ولی صفر Threat Actor، صفر Country و تنها یک Sector؛ گروه‌های MITRE در این دیتاست به‌صورت Intrusion Set ذخیره می‌شوند", "در حال رفع با دیتاست‌های مرجع (بخش ۶)"],
    ["۵", "تفسیر باریک NLQ", "«spearphishing» به تک‌تکنیک T1598.004 (فیشینگ صوتی) حل شد به‌جای خانوادهٔ گسترده‌تر (T1566/T1598)", "ثبت در بهبودهای بعدی پرامپت NLQ"],
    ["۶", "حل بی‌صدای نام ناموجود", "«Atlantis group» (که در پلتفرم نیست) به T1615 Group Policy Discovery گره خورد — به‌جای پیام «یافت نشد»", "ثبت: نیاز به آستانهٔ شباهت نام در resolve"],
])
p("نکتهٔ تاریخی مرتبط: در مسیر استریم LLM پیش‌تر باگ «nullnull» وجود داشت (دریافت content=null در چانک‌های gateway باعث الحاق رشتهٔ literal null می‌شد) که ریشه‌یابی و در هر دو مسیر Mistral و OpenAI در همین شاخه ترمیم شده بود؛ تمام تست‌های این چرخه نبود آن را تأیید کرد.")

h1("۶. وضعیت فعلی و گام‌های بعدی")
b([
    "واردکردن دیتاست مرجع رسمی OpenCTI (geography.json شامل کشورها/مناطق/شهرها و sectors.json) برای پرشدن صفحات جغرافیایی و سکتورها",
    "تبدیل و واردکردن کهکشان MISP (threat-actor.json) به STIX 2.1 — بازیگران تهدید واقعی همراه روابط targets به کشورها — برای پرمحتوا شدن تست UI گزارش‌های ۴ج/۴د",
    "گسترش اختیاری دکمهٔ «Generate AI report» به صفحات Intrusion Set (خانهٔ واقعی گروه‌ها در دیتاست ATT&CK)",
    "بهبودهای NLQ: آستانهٔ شباهت نام در resolve مقادیر و ترجیح خانوادهٔ گسترده‌تر تکنیک‌ها",
])
pb()

h1("ضمیمهٔ الف — وضعیت همهٔ mutationهای AI")
tb(["mutation", "کاربرد", "پیش از شاخه", "اکنون"], [
    ["aiFixSpelling / aiMakeShorter / aiMakeLonger / aiChangeTone / aiSummarize / aiExplain", "شش عمل متنی", "فعال (قفل EE در UI)", "فعال و تست‌شده"],
    ["aiNLQ", "جست‌وجوی زبان طبیعی", "فعال در بک‌اند؛ قفل بی‌صدای UI", "فعال و تست‌شده (گیت رفع شد)"],
    ["aiContainerGenerateReport", "گزارش کانتینر", "فعال (قفل EE)", "فعال و تست‌شده"],
    ["aiThreatGenerateReport", "گزارش بازیگر تهدید", "بدون resolver (null)", "پیاده‌سازی + دکمهٔ UI"],
    ["aiVictimGenerateReport", "گزارش قربانی (کشور/سکتور/منطقه)", "بدون resolver (null)", "پیاده‌سازی + دکمهٔ UI"],
    ["aiConvertIndicator", "تبدیل شاخص به STIX/SIGMA/YARA", "بدون resolver (null)", "پیاده‌سازی + دکمهٔ UI"],
    ["aiImproveWriting", "بهبود نگارش", "بدون resolver (null)", "پیاده‌سازی + آیتم منو"],
    ["aiTranslate", "ترجمهٔ وفادار", "وجود نداشت", "جدید: اسکیما + resolver + منو"],
    ["aiSummarizeFiles / aiConvertFilesToStix", "خلاصهٔ فایل / تبدیل فایل به STIX", "فعال (deprecated)", "بدون تغییر"],
])

h1("ضمیمهٔ ب — نقشهٔ مسیر فایل‌های کلیدی")
tb(["لایه", "فایل"], [
    ["فرانت — منوی Ask-AI", "src/private/components/common/form/TextFieldAskAI.tsx"],
    ["فرانت — کشو Insights و کش", "src/private/components/common/ai/AIInsights.tsx، AISummary*.tsx، src/utils/ai/insightsCache.ts"],
    ["فرانت — NLQ", "src/private/components/nav/TopBar.tsx + src/components/SearchInput.jsx"],
    ["فرانت — گزارش کانتینر", "src/private/components/common/stix_core_objects/StixCoreObjectFileExport.tsx + StixCoreObjectAskAI.tsx"],
    ["فرانت — دکمه‌های گزارش موجودیت", "src/private/components/common/ai/EntityAiReportButton.tsx (در Root صفحات Threat Actor، Country، Sector)"],
    ["فرانت — تبدیل Indicator", "src/private/components/observations/indicators/IndicatorDetails.tsx"],
    ["فرانت — پنل ایجنت", "src/private/components/agent/AgentPanel.tsx (مسیر /dashboard/agent)"],
    ["بک‌اند — اسکیما و resolver", "src/modules/ai/ai.graphql + ai-resolver.ts + ai-domain.ts"],
    ["بک‌اند — پروکسی ایجنت", "src/http/httpAgentProxy.ts"],
])

# ─────────────────────────── DOCX renderer ───────────────────────────
def _rtl_para(doc_par):
    pPr = doc_par._p.get_or_add_pPr()
    bidi = OxmlElement('w:bidi'); bidi.set(qn('w:val'), '1'); pPr.append(bidi)
    doc_par.alignment = WD_ALIGN_PARAGRAPH.RIGHT

def _style_run(run, size, bold=False, color=None, italic=False):
    run.font.name = 'Vazirmatn'
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color: run.font.color.rgb = RGBColor(*color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts'); rPr.append(rFonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rFonts.set(qn(attr), 'Vazirmatn')
    szCs = OxmlElement('w:szCs'); szCs.set(qn('w:val'), str(int(size * 2))); rPr.append(szCs)
    if bold:
        bCs = OxmlElement('w:bCs'); bCs.set(qn('w:val'), '1'); rPr.append(bCs)
    rtl = OxmlElement('w:rtl'); rtl.set(qn('w:val'), '1'); rPr.append(rtl)

def _cell_bg(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)

def build_docx(path):
    doc = Document()
    sec = doc.sections[0]
    sec.right_margin = sec.left_margin = Cm(2); sec.top_margin = sec.bottom_margin = Cm(2)
    # cover
    for _ in range(6): doc.add_paragraph()
    t = doc.add_paragraph(); _rtl_para(t); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _style_run(t.add_run(TITLE), 22, bold=True, color=ACCENT)
    s = doc.add_paragraph(); _rtl_para(s); s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _style_run(s.add_run(SUBTITLE), 13, color=GRAY)
    d = doc.add_paragraph(); _rtl_para(d); d.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _style_run(d.add_run(DATE_STR), 11, color=GRAY)
    doc.add_page_break()

    for block in C:
        kind = block[0]
        if kind == 'pagebreak':
            doc.add_page_break(); continue
        if kind == 'h1':
            par = doc.add_paragraph(); _rtl_para(par)
            par.paragraph_format.space_before = Pt(18); par.paragraph_format.space_after = Pt(8)
            _style_run(par.add_run(block[1]), 16, bold=True, color=ACCENT)
        elif kind == 'h2':
            par = doc.add_paragraph(); _rtl_para(par)
            par.paragraph_format.space_before = Pt(12); par.paragraph_format.space_after = Pt(6)
            _style_run(par.add_run(block[1]), 13, bold=True, color=DARK)
        elif kind == 'h3':
            par = doc.add_paragraph(); _rtl_para(par)
            _style_run(par.add_run(block[1]), 11.5, bold=True, color=DARK)
        elif kind == 'p':
            par = doc.add_paragraph(); _rtl_para(par)
            par.paragraph_format.space_after = Pt(6); par.paragraph_format.line_spacing = 1.15
            _style_run(par.add_run(block[1]), 10.5)
        elif kind == 'bullets':
            for item in block[1]:
                par = doc.add_paragraph(); _rtl_para(par)
                par.paragraph_format.space_after = Pt(3)
                _style_run(par.add_run('•  ' + item), 10.5)
        elif kind == 'table':
            header, rows = block[1], block[2]
            table = doc.add_table(rows=1 + len(rows), cols=len(header))
            table.style = 'Table Grid'
            tblPr = table._tbl.tblPr
            bidi = OxmlElement('w:bidiVisual'); tblPr.append(bidi)
            for j, htxt in enumerate(header):
                cell = table.rows[0].cells[j]
                _cell_bg(cell, 'EDE CFA'.replace(' ', ''))  # EDECFA
                cp = cell.paragraphs[0]; _rtl_para(cp)
                _style_run(cp.add_run(htxt), 10, bold=True, color=ACCENT)
            for i, row in enumerate(rows):
                for j, val in enumerate(row):
                    cell = table.rows[1 + i].cells[j]
                    cp = cell.paragraphs[0]; _rtl_para(cp)
                    _style_run(cp.add_run(val), 9.5)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
    doc.save(path)

# ─────────────────────────── PDF renderer ───────────────────────────
class ReportPDF(FPDF):
    def header(self):
        if self.page_no() > 1:
            self.set_font('Vazirmatn', '', 8); self.set_text_color(*GRAY)
            self.cell(0, 6, TITLE, align='C', new_x='LMARGIN', new_y='NEXT')
            self.ln(2)
    def footer(self):
        self.set_y(-14); self.set_font('Vazirmatn', '', 8); self.set_text_color(*GRAY)
        self.cell(0, 8, f"{self.page_no()}", align='C')

def build_pdf(path):
    pdf = ReportPDF('P', 'mm', 'A4')
    pdf.alias_nb_pages()
    pdf.add_font('Vazirmatn', '', FONT_PATH_M)
    pdf.add_font('Vazirmatn', 'B', FONT_PATH_B)
    pdf.set_text_shaping(True)
    pdf.set_auto_page_break(True, margin=18)
    pdf.set_margins(15, 15, 15)

    pdf.add_page()
    pdf.ln(55)
    pdf.set_font('Vazirmatn', 'B', 20); pdf.set_text_color(*ACCENT)
    pdf.multi_cell(0, 11, TITLE, align='C')
    pdf.ln(3)
    pdf.set_font('Vazirmatn', '', 13); pdf.set_text_color(*GRAY)
    pdf.multi_cell(0, 8, SUBTITLE, align='C')
    pdf.ln(2)
    pdf.set_font('Vazirmatn', '', 11)
    pdf.cell(0, 8, DATE_STR, align='C', new_x='LMARGIN', new_y='NEXT')
    pdf.add_page()

    for block in C:
        kind = block[0]
        if kind == 'pagebreak':
            pdf.add_page(); continue
        if kind == 'h1':
            pdf.ln(4); pdf.set_font('Vazirmatn', 'B', 15); pdf.set_text_color(*ACCENT)
            pdf.multi_cell(0, 9, block[1]); pdf.ln(1.5)
        elif kind == 'h2':
            pdf.ln(3); pdf.set_font('Vazirmatn', 'B', 12); pdf.set_text_color(*DARK)
            pdf.multi_cell(0, 8, block[1]); pdf.ln(1)
        elif kind == 'h3':
            pdf.set_font('Vazirmatn', 'B', 11); pdf.set_text_color(*DARK)
            pdf.multi_cell(0, 7, block[1])
        elif kind == 'p':
            pdf.set_font('Vazirmatn', '', 10); pdf.set_text_color(40, 44, 55)
            pdf.multi_cell(0, 6.2, block[1]); pdf.ln(1.5)
        elif kind == 'bullets':
            pdf.set_font('Vazirmatn', '', 10); pdf.set_text_color(40, 44, 55)
            for item in block[1]:
                pdf.multi_cell(0, 6.2, '•  ' + item); pdf.ln(0.6)
            pdf.ln(1)
        elif kind == 'table':
            header, rows = block[1], block[2]
            pdf.set_font('Vazirmatn', '', 8.6); pdf.set_text_color(40, 44, 55)
            pdf.set_draw_color(180, 175, 200); pdf.set_line_width(0.15)
            head_face = FontFace(family='Vazirmatn', emphasis='BOLD', color=ACCENT, fill_color=HDR_BG)
            with pdf.table(headings_style=head_face, line_height=5.4, padding=1.4,
                           col_widths=([1] * len(header))) as table:
                hr = table.row()
                for htxt in header: hr.cell(htxt)
                for row in rows:
                    r = table.row()
                    for val in row: r.cell(val)
            pdf.ln(3)
    pdf.output(path)

if __name__ == '__main__':
    base = '/home/saeed/opencti/ai-features-report/'
    docx_path = base + 'AI-Capabilities-Verified-Report.fa.docx'
    pdf_path = base + 'AI-Capabilities-Verified-Report.fa.pdf'
    build_docx(docx_path)
    build_pdf(pdf_path)
    print('DOCX:', docx_path)
    print('PDF :', pdf_path)
